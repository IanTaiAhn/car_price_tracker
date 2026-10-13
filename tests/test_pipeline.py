"""End-to-end check with synthetic alerts and a fake dealer page.  Run:  python -m pytest -q  (or python tests/test_pipeline.py)"""
from __future__ import annotations

import shutil
import sys
import tempfile
from datetime import date
from email.message import EmailMessage
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from cartracker import __main__ as app  # noqa: E402
from cartracker import dealer, export  # noqa: E402
from cartracker.config import load_config, match_vehicle  # noqa: E402


def card(lid, title, price, miles, old=None):
    old_html = f"<s>${old:,}</s> " if old else ""
    return (f'<table><tr><td><a href="https://click.ksl.com/r?u=https%3A%2F%2Fcars.ksl.com%2Flisting%2F{lid}">'
            f'{title}</a></td></tr><tr><td>{old_html}${price:,} · {miles:,} miles · Sandy, UT</td></tr></table>')


def eml(path, subject, date_hdr, cards, mid):
    m = EmailMessage()
    m["Subject"], m["Date"], m["Message-ID"], m["From"] = subject, date_hdr, mid, "alerts@ksl.com"
    m.set_content("html only")
    m.add_alternative("<html><body>" + "".join(cards) + "</body></html>", subtype="html")
    path.write_bytes(bytes(m))
    return str(path)


DEALER_HTML = """<html><head><script type="application/ld+json">
[{"@context":"https://schema.org","@type":"Car","name":"2024 Subaru Forester Touring",
  "vehicleIdentificationNumber":"JF2SKALC5RH400001","vehicleModelDate":"2024","brand":{"name":"Subaru"},
  "model":"Forester","vehicleConfiguration":"Touring","mileageFromOdometer":{"value":"18500","unitCode":"SMI"},
  "offers":{"@type":"Offer","price":"33995"},"url":"/used/2024-subaru-forester-1"},
 {"@context":"https://schema.org","@type":"Car","name":"2023 Toyota RAV4 Hybrid XLE Premium",
  "vehicleIdentificationNumber":"4T3RWRFV9PU000002","vehicleModelDate":"2023","brand":{"name":"Toyota"},
  "model":"RAV4 Hybrid","vehicleConfiguration":"XLE Premium","mileageFromOdometer":{"value":"21000"},
  "offers":{"price":"31500"},"url":"/used/2023-rav4-2"},
 {"@context":"https://schema.org","@type":"Car","name":"2021 Subaru Forester Premium",
  "vehicleIdentificationNumber":"JF2SKAJC0MH000003","vehicleModelDate":"2021","brand":{"name":"Subaru"},
  "model":"Forester","mileageFromOdometer":{"value":"40000"},"offers":{"price":"22000"},"url":"/used/3"}]
</script></head><body></body></html>"""


def test_filters():
    cfg = load_config(ROOT / "config.yaml")
    ok = lambda **k: match_vehicle(cfg, k)[0] is not None  # noqa: E731
    yr = date.today().year
    assert ok(make="Toyota", model="RAV4 Hybrid", trim="XLE Premium", year=2023, mileage=20000)
    assert not ok(make="Toyota", model="RAV4", trim="XLE", year=2023, mileage=20000, title="2023 Toyota RAV4 Hybrid XLE")
    assert not ok(make="Toyota", model="RAV4", trim="XLE Premium", year=2023, mileage=20000)  # gas
    assert not ok(make="Toyota", model="RAV4 Prime", trim="XSE", year=2023, mileage=20000, fuel="Plug-in Hybrid")
    assert not ok(make="Subaru", model="Forester Hybrid", trim="Touring", year=2025, mileage=5000)
    assert not ok(make="Subaru", model="Forester", trim="Premium", year=2022, mileage=20000)
    assert ok(make="Subaru", model="Forester", trim="Wilderness", year=2022, mileage=20000)
    assert ok(make="Mazda", model="CX-5", trim="2.5 S Carbon Edition", year=2025, mileage=9000)
    assert not ok(make="Mazda", model="CX-5", trim="2.5 Carbon Turbo", year=2025, mileage=9000)
    assert not ok(make="Mazda", model="CX-5", trim="Carbon Edition", year=2024, mileage=9000)
    cap2020 = 12000 * (yr - 2020 + 1)
    assert ok(make="Subaru", model="Forester", trim="Limited", year=2020, mileage=cap2020)
    assert not ok(make="Subaru", model="Forester", trim="Limited", year=2020, mileage=cap2020 + 1)


def test_end_to_end():
    tmp = Path(tempfile.mkdtemp())
    try:
        shutil.copy(ROOT / "config.yaml", tmp / "config.yaml")
        cfg_path = str(tmp / "config.yaml")
        e1 = eml(tmp / "a.eml", "New listings match your saved search", "Mon, 05 Oct 2026 09:00:00 -0600", [
            card(11111111, "2023 Toyota RAV4 Hybrid XLE Premium AWD", 32500, 21050),
            card(22222222, "2022 Subaru Forester Limited", 26900, 30000),
            card(33333333, "2022 Subaru Forester Premium", 24000, 30000),     # wrong trim
            card(44444444, "2025 Mazda CX-5 2.5 S Carbon Edition", 29900, 8000),
        ], "<a@ksl>")
        app.main(["--config", cfg_path, "--eml", e1, "--skip-dealers"])

        e2 = eml(tmp / "b.eml", "Price drop on a listing you favorited", "Thu, 08 Oct 2026 09:00:00 -0600",
                 [card(22222222, "2022 Subaru Forester Limited", 25500, 30000, old=26900)], "<b@ksl>")
        e3 = eml(tmp / "c.eml", "Your favorited listing is expiring soon", "Tue, 06 Oct 2026 09:00:00 -0600",
                 [card(44444444, "2025 Mazda CX-5 2.5 S Carbon Edition", 29900, 8000)], "<c@ksl>")

        # fake a dealer whose page carries JSON-LD; one of its cars is the same RAV4 seen on KSL (fuzzy dup)
        cfg = load_config(cfg_path)
        cfg.raw["dealers"] = [{"name": "Test Dealer", "enabled": True, "base_url": "https://dealer.test",
                               "inventory_urls": ["https://dealer.test/used/"]}]
        orig_scrape = dealer.DealerScraper.scrape
        dealer.DealerScraper.scrape = lambda self: dealer.parse_jsonld_vehicles(DEALER_HTML, "https://dealer.test/used/")
        orig_load = app.load_config
        app.load_config = lambda p: cfg
        try:
            app.main(["--config", cfg_path, "--eml", e2, e3])
            # second dealer pass: Forester sold
            dealer.DealerScraper.scrape = lambda self: [
                v for v in dealer.parse_jsonld_vehicles(DEALER_HTML, "https://dealer.test/used/")
                if v["vin"] != "JF2SKALC5RH400001"]
            app.main(["--config", cfg_path, "--skip-email"])
        finally:
            dealer.DealerScraper.scrape = orig_scrape
            app.load_config = orig_load

        from cartracker.db import DB
        db = DB(cfg.db_path)
        rows = {r["source_id"]: r for r in db.conn.execute("SELECT * FROM listings")}
        assert "33333333" not in rows, "wrong trim should be filtered"
        assert rows["22222222"]["price"] == 25500 and rows["22222222"]["first_price"] == 26900
        assert rows["44444444"]["status"] == "likely_gone"          # expiring alert >3 days ago, no news since
        assert rows["JF2SKALC5RH400001"]["status"] == "gone"        # missing from dealer's later scrape
        assert "JF2SKAJC0MH000003" not in rows                      # Premium trim filtered
        rav_dealer = rows["4T3RWRFV9PU000002"]
        assert rav_dealer["possible_dup_of"] == rows["11111111"]["id"], "KSL vs dealer RAV4 should be flagged"
        assert rows["11111111"]["trim"] == "XLE Premium"

        # notes survive a rebuild
        wb = load_workbook(cfg.xlsx_path)
        assert wb.sheetnames == ["RAV4 Hybrid", "Forester", "CX-5 Carbon", "Trends", "Price History"]
        ws = wb["Forester"]
        hdr = [c.value for c in ws[export.HEADER_ROW]]
        r = next(r for r in range(export.HEADER_ROW + 1, ws.max_row + 1)
                 if ws.cell(r, hdr.index("ID") + 1).value == rows["22222222"]["id"])
        ws.cell(r, hdr.index("Notes") + 1, "call seller")
        wb.save(cfg.xlsx_path)
        app.main(["--config", cfg_path, "--export-only"])
        assert export.read_notes(cfg.xlsx_path)[rows["22222222"]["id"]] == "call seller"
        print("xlsx:", cfg.xlsx_path)
        return cfg.xlsx_path
    finally:
        pass  # leave tmp for inspection


if __name__ == "__main__":
    test_filters()
    print("filters ok")
    test_end_to_end()
    print("end-to-end ok")
