"""Rebuild the spreadsheet from the database.

Tabs: one per tracked vehicle, then Trends and Price History.
The Notes column is yours: it's read back from the existing file before each rebuild, so it survives.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import Config
from .db import DB

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3A5F")
NOTES_FILL = PatternFill("solid", start_color="FFF2CC")
GONE_FONT = Font(name=FONT, color="999999")
GREEN = PatternFill("solid", start_color="D9EAD3")
RED = PatternFill("solid", start_color="F4CCCC")
LEGEND = ("Rebuilt every run from data/tracker.db. Only the yellow Notes column is yours to edit: it is kept "
          "across runs. Status: active / likely_gone (KSL, inferred: no alert for a while or expired) / gone "
          "(dealer, confirmed missing from its site). Green = price below first-seen; red = above.")

COLUMNS = [  # header, width
    ("ID", 6), ("Status", 11), ("Year", 6), ("Trim", 14), ("Price", 10), ("First Price", 11),
    ("Change $", 10), ("Change %", 9), ("Lowest", 10), ("Mileage", 9), ("Mile Cap", 9),
    ("Days Listed", 10), ("Source", 12), ("Also On", 22), ("Possible Dup", 12), ("Title", 38),
    ("Link", 8), ("VIN", 19), ("First Seen", 11), ("Last Seen", 11), ("Gone", 11), ("Notes", 30),
]
COL = {h: i + 1 for i, (h, _) in enumerate(COLUMNS)}
HEADER_ROW = 3


def _letter(h):
    return get_column_letter(COL[h])


def _d(s):
    return datetime.fromisoformat(s).date() if s else None


def read_notes(path: Path) -> dict[int, str]:
    if not path.exists():
        return {}
    notes = {}
    wb = load_workbook(path, read_only=True)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = None
        for row in rows:
            if row and "ID" in row and "Notes" in row:
                header = row
                break
        if not header:
            continue
        i_id, i_n = header.index("ID"), header.index("Notes")
        for row in rows:
            if row and len(row) > i_n and isinstance(row[i_id], int) and row[i_n]:
                notes[row[i_id]] = str(row[i_n])
    wb.close()
    return notes


def build(cfg: Config, db: DB) -> Path:
    path = cfg.xlsx_path
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    c = db.conn

    for spec in cfg.vehicles:
        ws = wb.create_sheet(spec.sheet[:31])
        _vehicle_sheet(cfg, c, ws, spec)
    _trends_sheet(cfg, c, wb.create_sheet("Trends"))
    _history_sheet(cfg, c, wb.create_sheet("Price History"))

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.font.name != FONT:
                    f = cell.font
                    cell.font = Font(name=FONT, bold=f.bold, color=f.color, size=f.size, italic=f.italic,
                                     underline=f.underline)
    wb.save(path)
    return path


def _header(ws, cols, row=HEADER_ROW):
    for i, (h, w) in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=3)


def _vehicle_sheet(cfg, c, ws, spec):
    ws["A1"] = f"{spec.sheet} — {spec.make} {spec.model}, {spec.min_year}–{spec.max_year or 'newest'}, " \
               f"trims: {', '.join(spec.trims)}"
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    ws["A2"] = LEGEND
    ws["A2"].font = Font(name=FONT, italic=True, size=9, color="666666")
    _header(ws, COLUMNS)

    rows = c.execute("""SELECT * FROM listings WHERE vehicle_key=? AND dup_of IS NULL
                        ORDER BY CASE status WHEN 'active' THEN 0 WHEN 'likely_gone' THEN 1 ELSE 2 END,
                                 price IS NULL, price""", (spec.key,)).fetchall()
    r = HEADER_ROW
    for row in rows:
        r += 1
        dups = c.execute("SELECT source, price, status FROM listings WHERE dup_of=?", (row["id"],)).fetchall()
        status = row["status"]
        if status != "active" and any(d["status"] == "active" for d in dups):
            status = "active"
        also = ", ".join(f"{d['source']} ${d['price']:,}" if d["price"] else d["source"] for d in dups)
        vals = {
            "ID": row["id"], "Status": status, "Year": row["year"], "Trim": row["trim"],
            "Price": row["price"], "First Price": row["first_price"], "Lowest": row["min_price"],
            "Mileage": row["mileage"], "Mile Cap": cfg.mileage_cap(row["year"]) if row["year"] else None,
            "Source": row["source"], "Also On": also or None,
            "Possible Dup": f"#{row['possible_dup_of']}" if row["possible_dup_of"] else None,
            "Title": row["title"], "VIN": row["vin"],
            "First Seen": _d(row["first_seen"]), "Last Seen": _d(row["last_seen"]),
            "Gone": _d(row["gone_at"]), "Notes": row["notes"],
        }
        for h, v in vals.items():
            ws.cell(row=r, column=COL[h], value=v)
        P, F, FS, G = (f"{_letter(h)}{r}" for h in ("Price", "First Price", "First Seen", "Gone"))
        ws.cell(row=r, column=COL["Change $"], value=f'=IF(AND(ISNUMBER({P}),ISNUMBER({F})),{P}-{F},"")')
        ws.cell(row=r, column=COL["Change %"],
                value=f'=IF(AND(ISNUMBER({P}),ISNUMBER({F}),{F}<>0),({P}-{F})/{F},"")')
        ws.cell(row=r, column=COL["Days Listed"], value=f'=IF(ISNUMBER({G}),{G},TODAY())-{FS}')
        if row["url"]:
            link = ws.cell(row=r, column=COL["Link"], value="open")
            link.hyperlink = row["url"]
            link.font = Font(name=FONT, color="0563C1", underline="single")
        ws.cell(row=r, column=COL["Notes"]).fill = NOTES_FILL
        if status != "active":
            for h in vals:
                if h != "Notes":
                    ws.cell(row=r, column=COL[h]).font = GONE_FONT

    last = max(r, HEADER_ROW + 1)
    for h, fmt in (("Price", "$#,##0"), ("First Price", "$#,##0"), ("Lowest", "$#,##0"),
                   ("Change $", "$#,##0;-$#,##0;-"), ("Change %", "0.0%;-0.0%;-"), ("Mileage", "#,##0"),
                   ("Mile Cap", "#,##0"), ("Days Listed", "0"), ("First Seen", "yyyy-mm-dd"),
                   ("Last Seen", "yyyy-mm-dd"), ("Gone", "yyyy-mm-dd")):
        for rr in range(HEADER_ROW + 1, last + 1):
            ws.cell(row=rr, column=COL[h]).number_format = fmt
    if r == HEADER_ROW:
        ws.cell(row=HEADER_ROW + 1, column=1, value="No matching listings yet.").font = Font(name=FONT, italic=True)
    ch = f"{_letter('Change $')}{HEADER_ROW + 1}:{_letter('Change $')}{last}"
    ws.conditional_formatting.add(ch, CellIsRule(operator="lessThan", formula=["0"], fill=GREEN))
    ws.conditional_formatting.add(ch, CellIsRule(operator="greaterThan", formula=["0"], fill=RED))
    pr = f"{_letter('Price')}{HEADER_ROW + 1}:{_letter('Price')}{last}"
    ws.conditional_formatting.add(
        pr, FormulaRule(formula=[f'AND(ISNUMBER({_letter("Price")}{HEADER_ROW + 1}),'
                                 f'{_letter("Price")}{HEADER_ROW + 1}<{_letter("First Price")}{HEADER_ROW + 1})'],
                        fill=GREEN))
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(COLUMNS))}{last}"
    ws.row_dimensions[HEADER_ROW].height = 30


def _trends_sheet(cfg, c, ws):
    ws["A1"] = "Inventory movement per run"
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    ws["A2"] = ("One row per run. New / Gone count changes since the previous run. Median and lowest are over "
                "active listings only (duplicates counted once).")
    ws["A2"].font = Font(name=FONT, italic=True, size=9, color="666666")
    cols = [("Run Date", 18)]
    for spec in cfg.vehicles:
        cols += [(f"{spec.sheet} Active", 10), (f"{spec.sheet} New", 9), (f"{spec.sheet} Gone", 9),
                 (f"{spec.sheet} Median", 11), (f"{spec.sheet} Lowest", 11)]
    _header(ws, cols)
    ws.freeze_panes = "B4"
    runs = c.execute("SELECT id, started_at FROM runs ORDER BY id").fetchall()
    r = HEADER_ROW
    for run in runs:
        r += 1
        ws.cell(row=r, column=1, value=datetime.fromisoformat(run["started_at"])).number_format = "yyyy-mm-dd hh:mm"
        stats = {s["vehicle_key"]: s for s in c.execute("SELECT * FROM run_stats WHERE run_id=?", (run["id"],))}
        for i, spec in enumerate(cfg.vehicles):
            s = stats.get(spec.key)
            base = 2 + i * 5
            for j, k in enumerate(("active", "new_count", "gone_count", "median_price", "min_price")):
                cell = ws.cell(row=r, column=base + j, value=s[k] if s else None)
                if k in ("median_price", "min_price"):
                    cell.number_format = "$#,##0"
    if r > HEADER_ROW + 1:
        chart = LineChart()
        chart.title = "Median asking price (active listings)"
        chart.y_axis.title = "Price ($)"
        chart.y_axis.numFmt = "$#,##0"
        chart.height, chart.width = 9, 22
        for i, spec in enumerate(cfg.vehicles):
            col = 2 + i * 5 + 3
            chart.add_data(Reference(ws, min_col=col, min_row=HEADER_ROW, max_row=r), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=HEADER_ROW + 1, max_row=r))
        ws.add_chart(chart, f"B{r + 3}")


def _history_sheet(cfg, c, ws):
    ws["A1"] = "Every observed event, newest first"
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    cols = [("Date", 18), ("Vehicle", 14), ("ID", 6), ("Year", 6), ("Trim", 14), ("Source", 12),
            ("Event", 12), ("Price", 10), ("Mileage", 9), ("Title", 40)]
    _header(ws, cols)
    sheet_by_key = {s.key: s.sheet for s in cfg.vehicles}
    r = HEADER_ROW
    for h in c.execute("""SELECT h.*, l.vehicle_key, l.year, l.trim, l.source, l.mileage, l.title
                          FROM price_history h JOIN listings l ON l.id=h.listing_id
                          ORDER BY h.observed_at DESC, h.rowid DESC"""):
        r += 1
        vals = [datetime.fromisoformat(h["observed_at"]), sheet_by_key.get(h["vehicle_key"], h["vehicle_key"]),
                h["listing_id"], h["year"], h["trim"], h["source"], h["event"], h["price"], h["mileage"], h["title"]]
        for i, v in enumerate(vals, 1):
            ws.cell(row=r, column=i, value=v)
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row=r, column=8).number_format = "$#,##0"
        ws.cell(row=r, column=9).number_format = "#,##0"
        ev = ws.cell(row=r, column=7)
        if h["event"] == "price_drop":
            ev.fill = GREEN
        elif h["event"] == "price_rise":
            ev.fill = RED
    ws.auto_filter.ref = f"A{HEADER_ROW}:J{max(r, HEADER_ROW + 1)}"
